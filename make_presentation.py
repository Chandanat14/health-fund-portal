# make_presentation.py
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor

def add_title_slide(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    title = slide.shapes.title
    subtitle = slide.placeholders[1]
    title.text = "Health Fund Allocation Portal"
    subtitle.text = "Blockchain-Powered Transparent Healthcare Claims System"

def add_bullet_slide(prs, title, bullets):
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    shapes = slide.shapes
    title_shape = shapes.title
    body_shape = shapes.placeholders[1]
    title_shape.text = title
    tf = body_shape.text_frame
    tf.clear()
    for bullet in bullets:
        p = tf.add_paragraph()
        p.text = bullet
        p.level = 0
        p.font.size = Pt(18)

# Create presentation
prs = Presentation()
prs.slide_width = Inches(13.33)
prs.slide_height = Inches(7.5)

# Slide 1: Title
add_title_slide(prs)

# Slide 2: Problem
add_bullet_slide(prs, "Problem Statement", [
    "Manual claim processing causes delays and errors",
    "High risk of fraud and duplicate claims",
    "No transparency in fund allocation",
    "Difficulty tracking family-wise annual limits"
])

# Slide 3: Solution
add_bullet_slide(prs, "Solution Overview", [
    "Web-based portal with role-based access",
    "Blockchain for immutable claim logging",
    "Auto-generated PDF certificates with QR verification",
    "Bulk citizen data upload via XLSX"
])

# Slide 4: Architecture
add_bullet_slide(prs, "System Architecture", [
    "Frontend: Flask + HTML/Jinja2",
    "Backend: SQLite (thread-safe)",
    "Blockchain: Ethereum (Ganache local chain)",
    "Web3.py for smart contract interaction",
    "ReportLab for PDF generation",
    "QR codes for certificate verification"
])

# Slide 5: User Roles
add_bullet_slide(prs, "User Roles", [
    "District Admin → Creates Socio-Economic & Treasurer accounts",
    "Socio-Economic Admin → Manages citizen data & news",
    "Treasurer → Registers hospitals, approves/rejects claims",
    "Hospital Staff → Submits claims with bill PDFs"
])

# Slide 6: Socio-Economic Admin
add_bullet_slide(prs, "Socio-Economic Admin", [
    "Add/edit citizen: Aadhaar, name, income, family",
    "Upload photo (JPG/PNG/PDF)",
    "Bulk upload via XLSX with validation",
    "View all registered citizens"
])

# Slide 7: Treasurer Dashboard
add_bullet_slide(prs, "Treasurer Dashboard", [
    "Register hospitals with public/private keys",
    "View pending claims with bill PDFs",
    "Approve/Reject with reason",
    "View hospital-wise and overall stats"
])

# Slide 8: Hospital Staff
add_bullet_slide(prs, "Hospital Staff Workflow", [
    "Enter patient ID, Aadhaar, operation, amount, place",
    "Upload bill PDF",
    "System checks family annual limit (₹5 Lakh)",
    "Submit → recorded on blockchain"
])

# Slide 9: Blockchain
add_bullet_slide(prs, "Blockchain Integration", [
    "Smart Contract Functions:",
    "   • registerHospital(address)",
    "   • submitClaim(hospital, amount)",
    "   • approveClaim(claimId)",
    "   • rejectClaim(claimId)",
    "Event: ClaimSubmitted(claimId, hospital, amount)"
])

# Slide 10: PDF Certificate
add_bullet_slide(prs, "PDF Certificate Generation", [
    "Generated automatically on approval",
    "Includes: Claim ID, Amount, Hospital, QR Code",
    "Watermark: 'Health Portal'",
    "Emblem of India + border",
    "Downloadable by Treasurer & Hospital"
])

# Slide 11: Security
add_bullet_slide(prs, "Security Features", [
    "Password hashing (Werkzeug)",
    "Session: HttpOnly, SameSite=Lax, 1-hour expiry",
    "File upload validation (extensions, secure_filename)",
    "Input sanitization & logging",
    "Thread-local SQLite connections"
])

# Slide 12: Key Features
add_bullet_slide(prs, "Key Features", [
    "Indian number formatting (Jinja filter)",
    "XLSX bulk import with Aadhaar validation",
    "Family-wise claim limit enforcement",
    "Re-upload bill for rejected claims",
    "Comprehensive logging"
])

# Slide 13: Demo Flow
add_bullet_slide(prs, "Live Demo Flow", [
    "1. Login as Hospital → Submit Claim",
    "2. Treasurer → Review & Approve",
    "3. Certificate Generated → Download",
    "4. Scan QR → Verify on Blockchain"
])

# Slide 14: Future Enhancements
add_bullet_slide(prs, "Future Enhancements", [
    "Public claim verification portal",
    "SMS/Email alerts on status change",
    "IPFS storage for bill PDFs",
    "Digital signature on certificates",
    "Mobile app for hospital staff"
])

# Slide 15: Thank You
slide = prs.slides.add_slide(prs.slide_layouts[1])
title = slide.shapes.title
content = slide.placeholders[1].text_frame
title.text = "Thank You!"
content.clear()
p = content.add_paragraph()
p.text = "Project by: [Your Name]"
p.font.size = Pt(24)
p.alignment = PP_ALIGN.CENTER
p = content.add_paragraph()
p.text = "GitHub: github.com/yourusername/health-fund-portal"
p.font.size = Pt(18)
p.alignment = PP_ALIGN.CENTER
p = content.add_paragraph()
p.text = "Live Demo: http://yourdomain.com"
p.font.size = Pt(18)
p.alignment = PP_ALIGN.CENTER

# Save
prs.save("Health_Fund_Portal_Presentation.pptx")
print("PPTX generated: Health_Fund_Portal_Presentation.pptx")